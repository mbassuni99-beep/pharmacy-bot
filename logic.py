import os
import datetime
import calendar
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Local import
try:
    from db import load_data, get_custody_info
except ImportError:
    from .db import load_data, get_custody_info

def fmt(val):
    return f"{float(val or 0):,.2f}"

def progress_bar(percent, length=12):
    percent = max(0.0, min(100.0, percent))
    filled = int(round((percent / 100.0) * length))
    return "█" * filled + "░" * (length - filled)

def get_pharmacist_title(user):
    first_name = getattr(user, 'first_name', '') or ""
    last_name = getattr(user, 'last_name', '') or ""
    full = f"{first_name} {last_name}".strip()
    if not full:
        full = getattr(user, 'username', 'Pharmacist') or "Pharmacist"
    if not full.lower().startswith("dr.") and not full.startswith("د."):
        full = f"Dr. {full}"
    return full

# -------------------------------------------------------------
# 1. Daily Sales Calculation
# -------------------------------------------------------------
def calculate_sales_metrics(daily_sales, cash_sales, wasfaty_today, pharmacist_name="Dr. Pharmacist", day_num=None):
    data = load_data()
    now = datetime.datetime.now()
    today = now.date()
    if day_num is None:
        day_num = today.day

    _, last_day = calendar.monthrange(today.year, today.month)
    days_in_month = last_day
    remaining_days = max(1, days_in_month - day_num)

    # Compute cumulative up to yesterday
    cum_total = data.get("base_total_cum", 0.0)
    cum_wasfaty = data.get("base_wasfaty_cum", 0.0)

    for d in range(1, day_num):
        d_str = str(d)
        if d_str in data.get("history", {}):
            cum_total += data["history"][d_str].get("daily", 0.0)
            cum_wasfaty += data["history"][d_str].get("wasfaty", 0.0)

    # Add today
    total_cum = cum_total + daily_sales
    wasfaty_cum = cum_wasfaty + wasfaty_today
    non_wasfaty_cum = max(0.0, total_cum - wasfaty_cum)

    target_total = data.get("target_total", 150000.0)
    target_non_wasfaty = data.get("target_non_wasfaty", 100000.0)

    # Expected Sales (Run-rate formula)
    exp_total = (total_cum / day_num) * days_in_month if day_num > 0 else 0.0
    exp_non_wasfaty = (non_wasfaty_cum / day_num) * days_in_month if day_num > 0 else 0.0

    # Achievement % on Expected
    total_ach = (exp_total / target_total) * 100 if target_total > 0 else 0.0
    nw_ach = (exp_non_wasfaty / target_non_wasfaty) * 100 if target_non_wasfaty > 0 else 0.0

    # Status Emojis
    emoji_tot = "🟢" if total_ach >= 100 else ("🟡" if total_ach >= 85 else "🔴")
    emoji_nw = "🟢" if nw_ach >= 100 else ("🟡" if nw_ach >= 85 else "🔴")

    # 100% Target Run-rate
    rem_100_tot = max(0.0, target_total - total_cum)
    req_100_tot = rem_100_tot / remaining_days
    if total_cum >= target_total:
        line_100_tot = "🎯 100% Target: Achieved! 🎉 (100% Complete)"
    else:
        line_100_tot = f"🎯 100% Target: {fmt(rem_100_tot)} left ({fmt(req_100_tot)}/day)"

    # 85% Incentives
    target_85_nw = target_non_wasfaty * 0.85
    rem_85_nw = max(0.0, target_85_nw - non_wasfaty_cum)
    req_85_nw = rem_85_nw / remaining_days

    target_85_tot = target_total * 0.85
    rem_85_tot = max(0.0, target_85_tot - total_cum)
    req_85_tot = rem_85_tot / remaining_days

    if non_wasfaty_cum >= target_85_nw:
        line_85_nw = "🎯 85% (Non-Wasfaty): Unlocked! 🎉 (Qualifying)"
    else:
        line_85_nw = f"🎯 85% (Non-Wasfaty): {fmt(rem_85_nw)} left ({fmt(req_85_nw)}/day)"

    if total_cum >= target_85_tot:
        line_85_tot = "🎯 85% (Total Sales): Unlocked! 🎉 (Qualifying)"
    else:
        line_85_tot = f"🎯 85% (Total Sales): {fmt(rem_85_tot)} left ({fmt(req_85_tot)}/day)"

    # FEATURE 4: Day-over-Day Growth vs Yesterday
    yesterday_num = day_num - 1
    y_str = str(yesterday_num)
    growth_line = ""
    growth_pct = 0.0
    if y_str in data.get("history", {}) and data["history"][y_str].get("daily", 0) > 0:
        y_daily = data["history"][y_str]["daily"]
        diff = daily_sales - y_daily
        growth_pct = (diff / y_daily) * 100
        if diff >= 0:
            growth_line = f"📊 vs Yesterday: +{growth_pct:.1f}% 📈 (+{fmt(diff)} SAR)"
        else:
            growth_line = f"📊 vs Yesterday: {growth_pct:.1f}% 📉 ({fmt(diff)} SAR)"

    # FEATURE 2: Staff Accountability & Time
    submit_time = now.strftime("%I:%M %p")
    footer_staff = f"👨‍⚕️ Submitted by: {pharmacist_name} ({submit_time})"

    # Compose final WhatsApp message
    msg_lines = [
        f"daiy sales: {fmt(daily_sales)}",
        f"cash sales: {fmt(cash_sales)}",
        f"expected total sales: {fmt(exp_total)}",
        f"expected non wasfaty sales: {fmt(exp_non_wasfaty)}",
        f"non wasfaty target: {fmt(target_non_wasfaty)}",
        f"total non wasfaty achivement: {nw_ach:.1f}% {emoji_nw}",
        f"total sales target: {fmt(target_total)}",
        f"total sales achivement: {total_ach:.1f}% {emoji_tot}",
        f"wasfaty sales till today: {fmt(wasfaty_cum)}",
        "━━━━━━━━━━━━━"
    ]

    if growth_line:
        msg_lines.append(growth_line)

    msg_lines.append(line_100_tot)
    msg_lines.append(line_85_nw)
    msg_lines.append(line_85_tot)
    msg_lines.append("━━━━━━━━━━━━━")
    msg_lines.append(footer_staff)

    whatsapp_msg = "\n".join(msg_lines)

    return {
        "daily_sales": daily_sales,
        "cash_sales": cash_sales,
        "wasfaty_today": wasfaty_today,
        "total_cum": total_cum,
        "wasfaty_cum": wasfaty_cum,
        "non_wasfaty_cum": non_wasfaty_cum,
        "exp_total": exp_total,
        "exp_non_wasfaty": exp_non_wasfaty,
        "total_ach": total_ach,
        "nw_ach": nw_ach,
        "emoji_tot": emoji_tot,
        "emoji_nw": emoji_nw,
        "target_total": target_total,
        "target_non_wasfaty": target_non_wasfaty,
        "rem_100_tot": rem_100_tot,
        "req_100_tot": req_100_tot,
        "rem_85_nw": rem_85_nw,
        "req_85_nw": req_85_nw,
        "rem_85_tot": rem_85_tot,
        "req_85_tot": req_85_tot,
        "growth_line": growth_line,
        "pharmacist_name": pharmacist_name,
        "submit_time": submit_time,
        "day_num": day_num,
        "days_in_month": days_in_month,
        "remaining_days": remaining_days,
        "whatsapp_msg": whatsapp_msg
    }

# -------------------------------------------------------------
# 2. Petty Cash Custody Audit Calculation
# -------------------------------------------------------------
def calculate_custody_audit(purchases_invoices, insurance_amounts, actual_cash, pharmacist_name="Dr. Pharmacist", notes=""):
    data = load_data()
    now = datetime.datetime.now()
    fixed = float(data.get("fixed_custody", 2000.0))
    invoices = float(purchases_invoices or 0.0)
    insurance = float(insurance_amounts or 0.0)
    actual = float(actual_cash or 0.0)

    total_docs = invoices + insurance
    expected = max(0.0, fixed - total_docs)
    diff = actual - expected

    if abs(diff) < 0.01:
        status_line = "✅ مطابقة تماماً (عجز: 0.00 SAR)"
        status_short = "مطابق"
        status_emoji = "🟢"
    elif diff < 0:
        status_line = f"🔴 عجز في العهدة: {fmt(abs(diff))} SAR ⚠️"
        status_short = f"عجز: {fmt(abs(diff))}"
        status_emoji = "🔴"
    else:
        status_line = f"🟢 زيادة في العهدة: +{fmt(diff)} SAR"
        status_short = f"زيادة: +{fmt(diff)}"
        status_emoji = "🟡"

    audit_time = now.strftime("%I:%M %p")
    audit_date = now.strftime("%d/%m/%Y")

    wa_lines = [
        "💼 *تقرير عهدة فرع زهرة الروضة 103* 💊",
        f"📅 *التاريخ:* {audit_date} ({audit_time})",
        "━━━━━━━━━━━━━━━━━",
        f"💵 *العهدة الأساسية للفرع:* {fmt(fixed)} SAR",
        f"🧾 *فواتير مشتريات ومصروفات:* {fmt(invoices)} SAR",
        f"📋 *مبالغ التأمين / الاسترداد:* {fmt(insurance)} SAR",
        f"📊 *إجمالي الفواتير والمستندات:* {fmt(total_docs)} SAR",
        "━━━━━━━━━━━━━━━━━",
        f"💰 *الكاش المفترض توفره بالدرج:* {fmt(expected)} SAR",
        f"💵 *الكاش الفعلي المعدود:* {fmt(actual)} SAR",
        f"⚖️ *نتيجة الجرد:* {status_line}",
        "━━━━━━━━━━━━━━━━━",
        f"👨‍⚕️ *جرد الصيدلي المسؤول:* {pharmacist_name}"
    ]

    if notes:
        wa_lines.insert(-1, f"📝 *ملاحظات:* {notes}")

    wa_custody_msg = "\n".join(wa_lines)

    return {
        "fixed_custody": fixed,
        "purchases_invoices": invoices,
        "insurance_amounts": insurance,
        "total_docs": total_docs,
        "expected_cash": expected,
        "actual_cash": actual,
        "diff": diff,
        "status_line": status_line,
        "status_short": status_short,
        "status_emoji": status_emoji,
        "pharmacist_name": pharmacist_name,
        "audit_time": audit_time,
        "audit_date": audit_date,
        "whatsapp_msg": wa_custody_msg
    }

# -------------------------------------------------------------
# 3. Comprehensive Excel Report Generation
# -------------------------------------------------------------
def generate_excel_report(output_path="/tmp/pharmacy_monthly_report.xlsx"):
    data = load_data()
    now = datetime.datetime.now()
    _, last_day = calendar.monthrange(now.year, now.month)

    wb = openpyxl.Workbook()
    
    # --- Sheet 1: Daily Sales ---
    ws_sales = wb.active
    ws_sales.title = f"Sales_{now.strftime('%b_%Y')}"
    ws_sales.views.sheetView[0].showGridLines = True

    navy_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_fill = PatternFill(start_color="2A4D69", end_color="2A4D69", fill_type="solid")
    alt_fill = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")
    white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D0D7DE'),
        right=Side(style='thin', color='D0D7DE'),
        top=Side(style='thin', color='D0D7DE'),
        bottom=Side(style='thin', color='D0D7DE')
    )

    ws_sales.merge_cells("A1:G1")
    t1 = ws_sales["A1"]
    t1.value = f"Zahrat Alrawda 103 - Daily Sales Report ({now.strftime('%B %Y')})"
    t1.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    t1.fill = navy_fill
    t1.alignment = center_align
    ws_sales.row_dimensions[1].height = 32

    sales_headers = ["Day", "Daily Sales", "Cash Sales", "Wasfaty Sales", "Non-Wasfaty Sales", "Cumulative Total", "Pharmacist & Notes"]
    ws_sales.append([])
    ws_sales.append(sales_headers)
    ws_sales.row_dimensions[3].height = 24

    for col_idx in range(1, len(sales_headers) + 1):
        c = ws_sales.cell(row=3, column=col_idx)
        c.font = white_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = thin_border

    cum_total = data.get("base_total_cum", 0.0)
    row_start = 4
    for day in range(1, last_day + 1):
        d_str = str(day)
        h = data.get("history", {}).get(d_str, {})
        d_val = h.get("daily", 0.0)
        c_val = h.get("cash", 0.0)
        w_val = h.get("wasfaty", 0.0)
        nw_val = max(0.0, d_val - w_val)
        pharm = h.get("pharmacist", "")
        time_rec = h.get("time", "")
        status_note = f"{pharm} ({time_rec})" if d_val > 0 else ""

        if d_val > 0:
            cum_total += d_val

        row_vals = [
            f"Day {day:02d}",
            d_val if d_val > 0 else "",
            c_val if c_val > 0 else "",
            w_val if w_val > 0 else "",
            nw_val if nw_val > 0 else "",
            cum_total if cum_total > 0 else "",
            status_note
        ]
        ws_sales.append(row_vals)
        cur_row = row_start + day - 1
        ws_sales.row_dimensions[cur_row].height = 20

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws_sales.cell(row=cur_row, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = center_align
            elif col_idx in [2, 3, 4, 5, 6]:
                cell.alignment = right_align
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = center_align

            if day % 2 == 0 and d_val == 0:
                cell.fill = alt_fill

    for col in ws_sales.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws_sales.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # --- Sheet 2: Petty Cash Custody Audit History ---
    ws_cust = wb.create_sheet(title="Petty_Cash_Custody")
    ws_cust.views.sheetView[0].showGridLines = True

    ws_cust.merge_cells("A1:H1")
    t2 = ws_cust["A1"]
    t2.value = f"Zahrat Alrawda 103 - Branch Petty Cash Float & Audit Log (Fixed: {fmt(data.get('fixed_custody', 2000))} SAR)"
    t2.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    t2.fill = navy_fill
    t2.alignment = center_align
    ws_cust.row_dimensions[1].height = 32

    cust_headers = ["Audit Date & Time", "Fixed Custody", "Purchases / فواتير", "Insurance / تأمين", "Expected Cash", "Actual Cash", "Audit Status", "Pharmacist"]
    ws_cust.append([])
    ws_cust.append(cust_headers)
    ws_cust.row_dimensions[3].height = 24

    for col_idx in range(1, len(cust_headers) + 1):
        c = ws_cust.cell(row=3, column=col_idx)
        c.font = white_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = thin_border

    history_records = data.get("custody_history", [])
    if not history_records:
        curr = get_custody_info()
        history_records = [{
            "date": now.strftime("%Y-%m-%d %I:%M %p"),
            "fixed": curr["fixed_custody"],
            "invoices": curr["purchases_invoices"],
            "insurance": curr["insurance_amounts"],
            "expected": curr["expected_cash"],
            "actual": curr["actual_cash"],
            "status": curr["status"],
            "pharmacist": curr["last_audit_pharmacist"] or "Current Balance"
        }]

    r_row = 4
    for rec in history_records:
        row_vals = [
            rec.get("date", ""),
            rec.get("fixed", 2000.0),
            rec.get("invoices", 0.0),
            rec.get("insurance", 0.0),
            rec.get("expected", 0.0),
            rec.get("actual", 0.0),
            rec.get("status", "مطابق"),
            rec.get("pharmacist", "")
        ]
        ws_cust.append(row_vals)
        ws_cust.row_dimensions[r_row].height = 20

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws_cust.cell(row=r_row, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = center_align
            elif col_idx in [2, 3, 4, 5, 6]:
                cell.alignment = right_align
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = center_align
        r_row += 1

    for col in ws_cust.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws_cust.column_dimensions[col_letter].width = max(max_len + 4, 15)

    wb.save(output_path)
    return output_path
