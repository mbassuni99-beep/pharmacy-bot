import os
import datetime
import calendar
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from db import load_data

def fmt(val):
    return f"{float(val or 0):,.2f}"

def progress_bar(percent, length=12):
    percent = max(0.0, min(100.0, percent))
    filled = int(round((percent / 100.0) * length))
    return "█" * filled + "░" * (length - filled)

def get_pharmacist_title(user):
    first_name = getattr(user, 'first_name', '') or ""
    last_name = getattr(user, 'last_name', '') or ""
    full = f"{first_name} {last_name}".strip()
    if not full:
        full = getattr(user, 'username', 'Pharmacist') or "Pharmacist"
    if not full.lower().startswith("dr.") and not full.startswith("د."):
        full = f"Dr. {full}"
    return full

def calculate_metrics(daily_sales, cash_sales, wasfaty_today, pharmacist_name="Dr. Pharmacist", day_num=None):
    data = load_data()
    now = datetime.datetime.now()
    today = now.date()
    if day_num is None:
        day_num = today.day

    _, last_day = calendar.monthrange(today.year, today.month)
    days_in_month = last_day
    remaining_days = max(1, days_in_month - day_num)

    # Compute cumulative up to yesterday
    cum_total = data.get("base_total_cum", 0.0)
    cum_wasfaty = data.get("base_wasfaty_cum", 0.0)

    for d in range(1, day_num):
        d_str = str(d)
        if d_str in data.get("history", {}):
            cum_total += data["history"][d_str].get("daily", 0.0)
            cum_wasfaty += data["history"][d_str].get("wasfaty", 0.0)

    # Add today
    total_cum = cum_total + daily_sales
    wasfaty_cum = cum_wasfaty + wasfaty_today
    non_wasfaty_cum = max(0.0, total_cum - wasfaty_cum)

    target_total = data.get("target_total", 150000.0)
    target_non_wasfaty = data.get("target_non_wasfaty", 100000.0)

    # Expected Sales (Run-rate formula)
    exp_total = (total_cum / day_num) * days_in_month if day_num > 0 else 0.0
    exp_non_wasfaty = (non_wasfaty_cum / day_num) * days_in_month if day_num > 0 else 0.0

    # Achievement % on Expected
    total_ach = (exp_total / target_total) * 100 if target_total > 0 else 0.0
    nw_ach = (exp_non_wasfaty / target_non_wasfaty) * 100 if target_non_wasfaty > 0 else 0.0

    # Status Emojis
    emoji_tot = "🟢" if total_ach >= 100 else ("🟡" if total_ach >= 85 else "🔴")
    emoji_nw = "🟢" if nw_ach >= 100 else ("🟡" if nw_ach >= 85 else "🔴")

    # 85% Incentives
    target_85_nw = target_non_wasfaty * 0.85
    rem_85_nw = max(0.0, target_85_nw - non_wasfaty_cum)
    req_85_nw = rem_85_nw / remaining_days

    target_85_tot = target_total * 0.85
    rem_85_tot = max(0.0, target_85_tot - total_cum)
    req_85_tot = rem_85_tot / remaining_days

    if non_wasfaty_cum >= target_85_nw:
        line_85_nw = "🎯 85% (Non-Wasfaty): Unlocked! 🎉 (Qualifying)"
    else:
        line_85_nw = f"🎯 85% (Non-Wasfaty): {fmt(rem_85_nw)} left ({fmt(req_85_nw)}/day)"

    if total_cum >= target_85_tot:
        line_85_tot = "🎯 85% (Total Sales): Unlocked! 🎉 (Qualifying)"
    else:
        line_85_tot = f"🎯 85% (Total Sales): {fmt(rem_85_tot)} left ({fmt(req_85_tot)}/day)"

    # FEATURE 4: Day-over-Day Growth vs Yesterday
    yesterday_num = day_num - 1
    y_str = str(yesterday_num)
    growth_line = ""
    growth_pct = 0.0
    if y_str in data.get("history", {}) and data["history"][y_str].get("daily", 0) > 0:
        y_daily = data["history"][y_str]["daily"]
        diff = daily_sales - y_daily
        growth_pct = (diff / y_daily) * 100
        if diff >= 0:
            growth_line = f"📊 vs Yesterday: +{growth_pct:.1f}% 📈 (+{fmt(diff)} SAR)"
        else:
            growth_line = f"📊 vs Yesterday: {growth_pct:.1f}% 📉 ({fmt(diff)} SAR)"

    # FEATURE 2: Staff Accountability & Time
    submit_time = now.strftime("%I:%M %p")
    footer_staff = f"👨‍⚕️ Submitted by: {pharmacist_name} ({submit_time})"

    # Compose final WhatsApp message
    msg_lines = [
        f"daiy sales: {fmt(daily_sales)}",
        f"cash sales: {fmt(cash_sales)}",
        f"expected total sales: {fmt(exp_total)}",
        f"expected non wasfaty sales: {fmt(exp_non_wasfaty)}",
        f"non wasfaty target: {fmt(target_non_wasfaty)}",
        f"total non wasfaty achivement: {nw_ach:.1f}% {emoji_nw}",
        f"total sales target: {fmt(target_total)}",
        f"total sales achivement: {total_ach:.1f}% {emoji_tot}",
        f"wasfaty sales till today: {fmt(wasfaty_cum)}",
        "━━━━━━━━━━━━━"
    ]

    if growth_line:
        msg_lines.append(growth_line)

    msg_lines.append(line_85_nw)
    msg_lines.append(line_85_tot)
    msg_lines.append("━━━━━━━━━━━━━")
    msg_lines.append(footer_staff)

    whatsapp_msg = "\n".join(msg_lines)

    return {
        "daily_sales": daily_sales,
        "cash_sales": cash_sales,
        "wasfaty_today": wasfaty_today,
        "total_cum": total_cum,
        "wasfaty_cum": wasfaty_cum,
        "non_wasfaty_cum": non_wasfaty_cum,
        "exp_total": exp_total,
        "exp_non_wasfaty": exp_non_wasfaty,
        "total_ach": total_ach,
        "nw_ach": nw_ach,
        "emoji_tot": emoji_tot,
        "emoji_nw": emoji_nw,
        "target_total": target_total,
        "target_non_wasfaty": target_non_wasfaty,
        "rem_85_nw": rem_85_nw,
        "req_85_nw": req_85_nw,
        "rem_85_tot": rem_85_tot,
        "req_85_tot": req_85_tot,
        "growth_line": growth_line,
        "pharmacist_name": pharmacist_name,
        "submit_time": submit_time,
        "day_num": day_num,
        "days_in_month": days_in_month,
        "remaining_days": remaining_days,
        "whatsapp_msg": whatsapp_msg
    }

def generate_excel_report(output_path="/tmp/pharmacy_monthly_report.xlsx"):
    data = load_data()
    now = datetime.datetime.now()
    _, last_day = calendar.monthrange(now.year, now.month)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Sales_{now.strftime('%b_%Y')}"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    navy_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_fill = PatternFill(start_color="2A4D69", end_color="2A4D69", fill_type="solid")
    alt_fill = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")
    white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D0D7DE'),
        right=Side(style='thin', color='D0D7DE'),
        top=Side(style='thin', color='D0D7DE'),
        bottom=Side(style='thin', color='D0D7DE')
    )

    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = f"Zahrat Alrawda 103 - Monthly Sales Report ({now.strftime('%B %Y')})"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = navy_fill
    title_cell.alignment = center_align
    ws.row_dimensions[1].height = 32

    headers = ["Day", "Daily Sales", "Cash Sales", "Wasfaty Sales", "Non-Wasfaty Sales", "Cumulative Total", "Status / Notes"]
    ws.append([])
    ws.append(headers)
    ws.row_dimensions[3].height = 24

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=3, column=col_idx)
        c.font = white_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = thin_border

    cum_total = 0.0
    row_start = 4
    for day in range(1, last_day + 1):
        d_str = str(day)
        h = data.get("history", {}).get(d_str, {})
        d_val = h.get("daily", 0.0)
        c_val = h.get("cash", 0.0)
        w_val = h.get("wasfaty", 0.0)
        nw_val = max(0.0, d_val - w_val)
        if d_val > 0:
            cum_total += d_val
        status_note = h.get("notes", "Submitted" if d_val > 0 else "")

        row_vals = [
            f"Day {day:02d}",
            d_val if d_val > 0 else "",
            c_val if c_val > 0 else "",
            w_val if w_val > 0 else "",
            nw_val if nw_val > 0 else "",
            cum_total if cum_total > 0 else "",
            status_note
        ]
        ws.append(row_vals)
        cur_row = row_start + day - 1
        ws.row_dimensions[cur_row].height = 20

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=cur_row, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = center_align
            elif col_idx in [2, 3, 4, 5, 6]:
                cell.alignment = right_align
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = center_align

            if day % 2 == 0 and d_val == 0:
                cell.fill = alt_fill

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    wb.save(output_path)
    return output_path
